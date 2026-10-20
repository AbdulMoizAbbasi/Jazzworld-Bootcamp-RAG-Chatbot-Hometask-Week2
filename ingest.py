
from pathlib import Path

from docx import Document
from dotenv import load_dotenv
from langchain_chroma import Chroma
from langchain_core.documents import Document as LangChainDocument
from langchain_google_genai import GoogleGenerativeAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
KNOWLEDGE_BASE_DIR = BASE_DIR / "knowledge_base"
VECTOR_DB_DIR = BASE_DIR / "chroma_db"

COLLECTION_NAME = "jazz_b2g_policies"


def load_policy_documents():
    documents = []

    for file_path in sorted(KNOWLEDGE_BASE_DIR.glob("*.docx")):
        doc = Document(file_path)

        paragraphs = [
            paragraph.text.strip()
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        ]

        # Extract text from tables too.
        table_text = []

        for table in doc.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    table_text.append(" | ".join(cells))

        text = "\n".join(paragraphs + table_text)

        if not text.strip():
            print(f"WARNING: No text extracted from {file_path.name}")
            continue

        documents.append(
            LangChainDocument(
                page_content=text,
                metadata={
                    "source": file_path.name,
                    "title": file_path.stem.replace("_", " "),
                },
            )
        )

    return documents


def main():
    # 1. Load the policy documents.
    documents = load_policy_documents()

    if not documents:
        raise ValueError(
            f"No DOCX documents found in {KNOWLEDGE_BASE_DIR}"
        )

    print(f"Documents loaded: {len(documents)}")

    # 2. Split documents into smaller chunks.
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=900,
        chunk_overlap=150,
        add_start_index=True,
    )

    chunks = splitter.split_documents(documents)

    print(f"Chunks created: {len(chunks)}")

    # 3. Configure Gemini embeddings.
    embeddings = GoogleGenerativeAIEmbeddings(
        model="models/gemini-embedding-001",
        output_dimensionality=768,
    )

    # 4. Create or update the persistent ChromaDB collection.
    vector_store = Chroma(
        collection_name=COLLECTION_NAME,
        embedding_function=embeddings,
        persist_directory=str(VECTOR_DB_DIR),
    )

    # Use deterministic IDs so rerunning ingestion updates
    # existing chunks instead of duplicating them.
    ids = []

    for chunk in chunks:
        source = chunk.metadata["source"]
        start_index = chunk.metadata["start_index"]
        ids.append(f"{source}::{start_index}")

    vector_store.add_documents(
        documents=chunks,
        ids=ids,
    )

    print("Ingestion completed successfully.")
    print(f"Documents processed: {len(documents)}")
    print(f"Chunks stored: {len(chunks)}")
    print(f"Vector database: {VECTOR_DB_DIR}")


if __name__ == "__main__":
    main()
